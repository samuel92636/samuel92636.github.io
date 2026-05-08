import sys
import subprocess

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import docx
except ImportError:
    install('python-docx')
    import docx

try:
    import pptx
except ImportError:
    install('python-pptx')
    import pptx

def extract_pptx(path, out_path):
    try:
        prs = pptx.Presentation(path)
        with open(out_path, 'w', encoding='utf-8') as f:
            for i, slide in enumerate(prs.slides):
                f.write(f"--- Slide {i+1} ---\n")
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        f.write(shape.text + "\n")
                f.write("\n")
        print(f"Successfully extracted {path} to {out_path}")
    except Exception as e:
        print(f"Error extracting {path}: {e}")

def extract_docx(path, out_path):
    try:
        doc = docx.Document(path)
        with open(out_path, 'w', encoding='utf-8') as f:
            for para in doc.paragraphs:
                f.write(para.text + "\n")
        print(f"Successfully extracted {path} to {out_path}")
    except Exception as e:
        print(f"Error extracting {path}: {e}")

if __name__ == "__main__":
    extract_pptx("大安高工-謝翔宇.pptx", "pptx_text.txt")
    extract_docx("申學與備審.docx", "docx_text.txt")
